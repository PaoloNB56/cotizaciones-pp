"""Manual PP feed update; never starts HTTP or writes the original quote files."""
from __future__ import annotations

import argparse
from contextlib import contextmanager
from datetime import date, datetime, timedelta, timezone
import hashlib
from io import BytesIO
import json
import math
import os
from pathlib import Path
import re
import sys
import time
from urllib.parse import urlparse
import uuid

import requests
from openpyxl import load_workbook
import update_bonos as legacy

ROOT = Path(__file__).resolve().parent
# Explicit publication list: never glob the project or upload source documents.
SEEDS = {f"{s}.json": f"{s}.json" for s in (
    "AL35", "AE38", "AL41", "AN29", "AO28", "S30N6", "X29Y6",
    "CCL", "MEP", "T13F6", "T15D5", "T30J6", "TTJ26", "TTM26",
)}
SEEDS.update({f"{s}.json": f"FCI/{s}.json" for s in ("BCACCA", "BCAHA", "BCMMA")})
EXPECTED_TICKERS = {s: s for s in legacy.INSTRUMENTOS_QT}
EXPECTED_TICKERS.update(AN29="AN29D", AO28="AO28D")
PROBE = "_pp_feed_probe.json"
def TODAY():
    return datetime.now(timezone(timedelta(hours=-3))).date()


class FeedError(Exception):
    """Only controlled messages belong in this exception (no response bodies/URLs)."""


def validate(data, label="serie", today=None):
    today = today or TODAY()
    if not isinstance(data, list) or not data:
        raise FeedError(f"{label}: serie vacía o formato distinto de una lista.")
    previous = ""
    for row in data:
        if not isinstance(row, dict) or set(row) != {"date", "close"}:
            raise FeedError(f"{label}: se esperan únicamente date y close.")
        ds, value = row["date"], row["close"]
        try:
            parsed = date.fromisoformat(ds)
        except (TypeError, ValueError):
            raise FeedError(f"{label}: fecha inválida.") from None
        if parsed.isoformat() != ds or parsed > today or ds <= previous:
            raise FeedError(f"{label}: fecha futura, duplicada o fuera de orden.")
        if type(value) not in (int, float) or not math.isfinite(value) or value <= 0:
            raise FeedError(f"{label}: precio inválido, no finito o no positivo.")
        previous = ds
    return data


def decode(body, label):
    try:
        data = json.loads(body.decode("utf-8-sig"))
    except (ValueError, UnicodeError):
        raise FeedError(f"{label}: no se recibió JSON válido.") from None
    return validate(data, label)


def encode(data):
    return (json.dumps(validate(data), ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode()


def digest(body):
    return hashlib.sha256(body).hexdigest()


def require_coverage(old, new, label):
    missing = {r["date"] for r in old} - {r["date"] for r in new}
    if missing:
        raise FeedError(f"{label}: faltan {len(missing)} fechas del histórico previo; se cancela el lote.")


def merge(old, new):
    merged = {r["date"]: r for r in old}
    merged.update({r["date"]: r for r in new})
    return validate([merged[d] for d in sorted(merged)])


def local_seeds():
    return {name: decode((ROOT / path).read_bytes(), name) for name, path in SEEDS.items()}


def quicktrade(symbol, old, today):
    cfg = legacy.INSTRUMENTOS_QT[symbol]
    # Use certificate verification in the new runner, including REQUESTS_CA_BUNDLE.
    legacy.VERIFY_SSL_QT = True
    body = legacy.descargar_contenido_quicktrade(
        cfg["simbolo"], cfg["fecha_desde"], today.strftime("%d/%m/%Y"))
    rows = parse_quicktrade(body, symbol, today)
    lower = datetime.strptime(cfg["fecha_desde"], "%d/%m/%Y").date().isoformat()
    if rows[0]["date"] < lower:
        raise FeedError(f"{symbol}: el proveedor respondió fuera del rango solicitado.")
    require_coverage([r for r in old if r["date"] >= lower], rows, symbol)
    return merge(old, rows)


def parse_quicktrade(body, symbol, today):
    book = load_workbook(BytesIO(body), read_only=True, data_only=True)
    try:
        source = book.worksheets[0].iter_rows(values_only=True)
        columns = [str(c).strip().lstrip("\ufeff") for c in next(source)]
        ticker_cols = [c for c in columns if c.startswith("S") and c.endswith("mbolo")]
        if len(set(columns)) != len(columns) or len(ticker_cols) != 1 or not {"Fecha", "Cierre"}.issubset(columns):
            raise FeedError(f"{symbol}: columnas de QuickTrade inesperadas.")
        rows = []
        for values in source:
            item = dict(zip(columns, values))
            if str(item[ticker_cols[0]]).strip() != EXPECTED_TICKERS[symbol]:
                raise FeedError(f"{symbol}: ticker del proveedor inesperado o ausente.")
            if item["Cierre"] is None or isinstance(item["Cierre"], bool):
                raise FeedError(f"{symbol}: cierre ausente o inválido; no se descartan filas.")
            rows.append(legacy.cotizacion_quicktrade_a_pp(item["Fecha"], item["Cierre"]))
        return validate(sorted(rows, key=lambda r: r["date"]), symbol, today)
    finally:
        book.close()


def dollars(symbol, old, today):
    cfg = legacy.DOLARAZO_DOLARES_CONFIG[symbol]
    start = date.fromisoformat(cfg["fecha_desde_iso"])
    payload = legacy.descargar_json_dolarazo(f"historicas/{cfg['casa']}")
    if not isinstance(payload, dict) or payload.get("ok") is not True or not isinstance(payload.get("data"), list) or not payload["data"]:
        raise FeedError(f"{symbol}: respuesta histórica de Dolarazo inválida o vacía.")
    # Check every source row before reusing the legacy parser, which skips bad rows.
    raw = []
    for row in payload["data"]:
        if not isinstance(row, dict) or type(row.get("venta")) not in (int, float):
            raise FeedError(f"{symbol}: venta ausente o inválida en Dolarazo.")
        raw.append({"date": row.get("fecha"), "close": row["venta"]})
    validate(sorted(raw, key=lambda r: str(r["date"])), symbol, today)
    history = validate(legacy.parsear_registros_dolarazo_historico(payload, start, today), symbol, today)
    current = legacy.descargar_json_dolarazo(f"dolares/{cfg['casa']}")
    if not isinstance(current, dict) or current.get("ok") is not True or not isinstance(current.get("data"), dict):
        raise FeedError(f"{symbol}: respuesta actual inválida.")
    item = current["data"]
    try:
        ds = legacy.parsear_fecha_actualizacion(str(item["fechaActualizacion"])).isoformat()
    except (KeyError, ValueError, TypeError):
        raise FeedError(f"{symbol}: fecha actual inválida.") from None
    validate([{"date": ds, "close": item.get("venta")}], symbol, today)
    latest = legacy.parsear_registro_dolarazo_actual(current, start, today)
    fetched = merge(history, latest)
    # Preserve all old values; refresh the last 7 days, including intraday corrections.
    overlap = date.fromisoformat(old[-1]["date"]) - timedelta(days=7)
    require_coverage([r for r in old if r["date"] >= overlap.isoformat()], fetched, symbol)
    # Dolarazo reports calendar days. A missing weekday in the new suffix is suspect.
    available = {r["date"] for r in fetched}
    day = date.fromisoformat(old[-1]["date"]) + timedelta(days=1)
    end = date.fromisoformat(fetched[-1]["date"])
    while day <= end:
        if day.weekday() < 5 and day.isoformat() not in available:
            raise FeedError(f"{symbol}: hueco entre el histórico previo y los precios nuevos.")
        day += timedelta(days=1)
    return merge(old, [r for r in fetched if r["date"] >= overlap.isoformat()])


def build(previous):
    result = dict(previous)
    today = TODAY()
    for symbol in legacy.INSTRUMENTOS_QT:
        name = f"{symbol}.json"
        if name not in SEEDS:
            raise FeedError("Instrumento nuevo: actualizar la lista explícita de publicación primero.")
        result[name] = quicktrade(symbol, previous[name], today)
    for symbol in legacy.DOLARAZO_DOLARES_CONFIG:
        result[f"{symbol}.json"] = dollars(symbol, previous[f"{symbol}.json"], today)
    return result


def download_bytes(url):
    # Separate request: never send the API credential to a CDN/public URL.
    try:
        response = requests.get(url, timeout=(10, 40), allow_redirects=False)
        if response.status_code != 200:
            raise FeedError(f"Descarga JSON: HTTP {response.status_code} (se requiere respuesta directa 200).")
        return response.content, response.headers
    except requests.RequestException:
        raise FeedError("Descarga JSON: error de red o certificado.") from None


class PCloud:
    def __init__(self):
        self.host = os.environ.get("PCLOUD_API_HOST", "")
        self.token = os.environ.get("PCLOUD_ACCESS_TOKEN", "")
        self.base = os.environ.get("PCLOUD_PUBLIC_BASE_URL", "").rstrip("/") + "/"
        try:
            self.folder = int(os.environ.get("PCLOUD_FOLDER_ID", ""))
        except ValueError:
            raise FeedError("Falta PCLOUD_FOLDER_ID numérico.") from None
        u = urlparse(self.base)
        if self.host not in {"api.pcloud.com", "eapi.pcloud.com"} or not self.token or self.folder <= 0:
            raise FeedError("Revisar PCLOUD_API_HOST, PCLOUD_ACCESS_TOKEN y PCLOUD_FOLDER_ID (no se permite la raíz).")
        if u.scheme != "https" or not u.hostname or u.query or u.fragment or u.username or u.password or u.port or u.path in {"", "/"}:
            raise FeedError("PCLOUD_PUBLIC_BASE_URL debe ser el enlace HTTPS directo de la subcarpeta pública, sin parámetros.")

    def call(self, method, **kwargs):
        data = {"access_token": self.token, **kwargs.pop("data", {})}
        try:
            response = requests.post(f"https://{self.host}/{method}", data=data,
                                     timeout=(10, 60), allow_redirects=False, **kwargs)
            if response.status_code != 200:
                raise FeedError(f"pCloud {method}: HTTP {response.status_code}.")
            payload = response.json()
        except (requests.RequestException, ValueError):
            raise FeedError(f"pCloud {method}: error de red, certificado o JSON.") from None
        if not isinstance(payload, dict) or type(payload.get("result")) is not int:
            raise FeedError(f"pCloud {method}: respuesta inválida.")
        if payload["result"] != 0:
            raise FeedError(f"pCloud {method}: código {payload['result']}; revisar región, permisos, token o cuota.")
        return payload

    def index(self):
        metadata = self.call("listfolder", data={"folderid": self.folder})["metadata"]
        if metadata.get("folderid") != self.folder or not isinstance(metadata.get("contents"), list):
            raise FeedError("pCloud: la carpeta no coincide o su listado es inválido.")
        return {m["name"]: m for m in metadata["contents"]}

    def read(self, name):
        item = self.index().get(name)
        if item is None:
            return None
        if item.get("isfolder") or type(item.get("fileid")) is not int:
            raise FeedError(f"{name}: se esperaba un archivo remoto.")
        link = self.call("getfilelink", data={"fileid": item["fileid"]})
        host = link["hosts"][0]
        path = link["path"]
        if not isinstance(host, str) or not host.endswith(".pcloud.com") or any(c in host for c in "/:@") or not isinstance(path, str) or not path.startswith("/"):
            raise FeedError("pCloud: enlace de descarga inesperado.")
        body, _ = download_bytes(f"https://{host}{path}")
        decode(body, name)
        return body

    def upload(self, name, body):
        if name not in SEEDS and name != PROBE:
            raise FeedError("Archivo fuera de la lista explícita de publicación.")
        decode(body, name)
        payload = self.call("uploadfile", data={"folderid": self.folder, "nopartial": 1},
                            files={"file": (name, body, "application/json")})
        metadata = payload.get("metadata", [])
        if len(metadata) != 1 or metadata[0].get("name") != name or metadata[0].get("parentfolderid") != self.folder or metadata[0].get("size") != len(body):
            raise FeedError(f"{name}: pCloud no confirmó el archivo completo en el destino esperado.")
        if self.read(name) != body:
            raise FeedError(f"{name}: lectura por API distinta del contenido enviado.")

    def public_check(self, name, body, attempts=7):
        wanted = decode(body, name)
        for attempt in range(attempts):
            try:
                received, headers = download_bytes(self.base + name)
                if decode(received, name) == wanted:
                    age = headers.get("Age", "")
                    max_age = re.search(r"(?:s-maxage|max-age)=(\d+)", headers.get("Cache-Control", ""))
                    return {"attempts": attempt + 1, "age_seconds": int(age) if age.isdigit() else None,
                            "max_age_seconds": int(max_age[1]) if max_age else None,
                            "etag_present": "ETag" in headers, "last_modified_present": "Last-Modified" in headers}
            except FeedError:
                pass
            if attempt + 1 < attempts:
                time.sleep(5)
        # A nonce is diagnostic only; passing it must not turn the stable-URL test green.
        try:
            received, _ = download_bytes(self.base + name + "?probe=" + uuid.uuid4().hex)
            fresh = decode(received, name) == wanted
        except FeedError:
            fresh = False
        raise FeedError(f"{name}: la URL fija no entrega la versión esperada; con parámetro nuevo coincide={fresh}. Revisar caché o enlace.")


def probe(client):
    # Two distinct values per invocation avoid an old cached probe giving a false pass.
    value = uuid.uuid4().int % 1_000_000_000 + 1
    for n in (value, value + 1):
        body = encode([{"date": TODAY().isoformat(), "close": n}])
        client.upload(PROBE, body)
        headers = client.public_check(PROBE, body)
    print("Prueba pCloud OK: dos versiones leídas en la misma URL. Cabeceras presentes:", headers)


def remote_previous(client, seeds, bootstrap=False):
    snapshots, previous = {}, {}
    for name in SEEDS:
        body = client.read(name)
        snapshots[name] = body
        if body is None:
            if not bootstrap:
                raise FeedError(f"{name}: falta el histórico remoto. No se reconstruye silenciosamente; usar bootstrap sólo para la primera carga.")
            previous[name] = seeds[name]
        else:
            remote = decode(body, name)
            require_coverage(seeds[name], remote, name)
            previous[name] = remote
    return previous, snapshots


def publish(client, candidates, snapshots, work):
    # Verify every existing public URL and detect a competing writer before any quote upload.
    for name in SEEDS:
        if client.read(name) != snapshots[name]:
            raise FeedError(f"{name}: cambió remotamente durante la preparación; volver a ejecutar.")
        if snapshots[name] is not None:
            client.public_check(name, snapshots[name])
    attempted = []
    try:
        for name in SEEDS:
            body = encode(candidates[name])
            if snapshots[name] is not None and decode(snapshots[name], name) == candidates[name]:
                continue
            attempted.append(name)  # Include uncertain network failures in recovery.
            client.upload(name, body)
            client.public_check(name, body)
            print(f"Publicado y verificado: {name}")
    except Exception:
        failed, new_files = [], []
        for name in reversed(attempted):
            if snapshots[name] is None:
                new_files.append(name)  # No destructive cleanup of a first publication.
                continue
            try:
                client.upload(name, snapshots[name])
                client.public_check(name, snapshots[name])
            except Exception:
                failed.append(name)
        recovery = {"restoration_failed": failed, "new_files_left": new_files, "attempted": attempted}
        (work / "recovery.json").write_text(json.dumps(recovery, indent=2), encoding="utf-8")
        print("Recuperación:", json.dumps(recovery))
        if failed or new_files:
            raise FeedError("Publicación parcial: revisar recovery.json y copias previas antes de actualizar PP.") from None
        raise FeedError("Falló la publicación; se restauraron y verificaron los archivos previos afectados.") from None
    return attempted


def execute(mode, work):
    seeds = local_seeds()
    client = None if mode == "local-check" else PCloud()
    if mode == "probe":
        probe(client)
        return {"status": "probe_verified", "files": []}
    if client is None:
        previous, snapshots = seeds, {}
    else:
        previous, snapshots = remote_previous(client, seeds, bootstrap=mode == "bootstrap")
    candidates = previous if mode == "bootstrap" else build(previous)
    report = []
    for name, rows in candidates.items():
        validate(rows, name)
        require_coverage(previous[name], rows, name)
        (work / name).write_bytes(encode(rows))
        if snapshots.get(name) is not None:
            backup = work / "previous"
            backup.mkdir(exist_ok=True)
            (backup / name).write_bytes(snapshots[name])
        kind = "descargado" if name[:-5] in legacy.INSTRUMENTOS_QT or name[:-5] in legacy.DOLARAZO_DOLARES_CONFIG else "conservado sin descarga"
        if mode == "bootstrap":
            kind = "histórico inicial conservado"
        item = {"file": name, "count": len(rows), "first": rows[0]["date"], "last": rows[-1]["date"], "source": kind, "sha256": digest(encode(rows))}
        report.append(item)
        print(f"{name}: {len(rows)} registros, {item['first']} a {item['last']} ({kind})")
    (work / "validation.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    if mode in {"update", "bootstrap"}:
        probe(client)
        published = publish(client, candidates, snapshots, work)
        return {"status": "published_verified", "published": published, "files": report}
    return {"status": "validated_not_published", "files": report}


@contextmanager
def single_local_writer():
    """OS lock releases on crash; serializes Termux/local launchers, not different devices."""
    directory = ROOT / ".cloud-work"
    directory.mkdir(exist_ok=True)
    with open(directory / "run.lock", "a+b") as stream:
        stream.seek(0)
        stream.write(b"0")
        stream.flush()
        stream.seek(0)
        try:
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(stream.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(stream, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError:
            raise FeedError("Ya hay una actualización en curso en este dispositivo.") from None
        try:
            yield
        finally:
            stream.seek(0)
            if os.name == "nt":
                msvcrt.locking(stream.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(stream, fcntl.LOCK_UN)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("mode", choices=["local-check", "check", "probe", "bootstrap", "update"], nargs="?", default="local-check")
    args = parser.parse_args(argv)
    work = ROOT / ".cloud-work" / (datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid.uuid4().hex[:8])
    work.mkdir(parents=True)
    try:
        with single_local_writer():
            report = execute(args.mode, work)
        code = 0
    except Exception as exc:
        # Never echo arbitrary exceptions: requests may embed signed URLs or credentials.
        message = str(exc) if isinstance(exc, FeedError) else f"Error {type(exc).__name__}; revisar proveedor, dependencias y configuración."
        report, code = {"status": "failed", "error": message}, 1
        print("ERROR:", message)
    (work / "result.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    summary = f"Resultado: {report['status']}\n"
    if code:
        summary += report["error"] + "\n"
    for item in report.get("files", []):
        summary += f"- {item['file']}: {item['count']} precios, última fecha {item['last']}; {item['source']}\n"
    if os.environ.get("GITHUB_STEP_SUMMARY"):
        with open(os.environ["GITHUB_STEP_SUMMARY"], "a", encoding="utf-8") as stream:
            stream.write(summary)
    print(f"{summary}Informe local: {work / 'result.json'}")
    return code


if __name__ == "__main__":
    raise SystemExit(main())
